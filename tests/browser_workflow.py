"""End-to-end journal workflow using a disposable database and real Chromium."""
import base64
import contextlib
import io
import json
import os
from pathlib import Path
import re
import socket
import urllib.request
import subprocess
import sys
import tempfile
import time
import zipfile

from docx import Document
from PIL import Image
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]
QA = ROOT / '.qa'
QA.mkdir(exist_ok=True)
PASSWORD = 'local-test-only-2026'


class Server:
    def __init__(self, data):
        self.data = data
        self.proc = None
    def start(self):
        self.log = open(Path(self.data) / 'browser-server.log', 'w+', encoding='utf-8')
        env = dict(os.environ, AIDITOR_DATA_DIR=self.data, PYTHONDONTWRITEBYTECODE='1', PYTHONUNBUFFERED='1')
        packaged = os.environ.get('AIDITOR_TEST_APP')
        command = [packaged] if packaged else [sys.executable, 'app.py']
        port = '0'
        if packaged:
            with socket.socket() as probe:
                probe.bind(('127.0.0.1',0));port=str(probe.getsockname()[1])
            self.url='http://127.0.0.1:'+port
        self.proc = subprocess.Popen(command + ['--no-browser', '--port', port], cwd=ROOT,
                                     env=env, stdout=self.log, stderr=subprocess.STDOUT)
        deadline = time.monotonic()+(90 if packaged else 20)
        while time.monotonic() < deadline:
            if packaged:
                try:
                    with urllib.request.urlopen(self.url+'/health',timeout=1) as response:
                        health=json.load(response)
                    if health.get('app')=='AI-ditor Plus' and health.get('version')=='2.0.2':return self.url
                except (OSError,ValueError):pass
            self.log.seek(0)
            content=self.log.read()
            match=re.search(r'http://127\.0\.0\.1:(\d+)', content)
            if match:
                self.url=match.group(0); return self.url
            if self.proc.poll() is not None:
                raise RuntimeError(content)
            time.sleep(.1)
        raise RuntimeError('Local server did not start')
    def stop(self):
        if self.proc:
            self.proc.terminate()
            try:self.proc.wait(timeout=10)
            except subprocess.TimeoutExpired:self.proc.kill();self.proc.wait(timeout=5)
            self.log.close();self.proc=None


def ready(page):
    page.wait_for_function('window.journalWorkspace?.ready && window.articleLibrary?.ready')
    page.wait_for_function('!document.getElementById("article-editor").inert')


def saved(page):
    page.wait_for_function('window.articleLibrary?.currentId && !articleLibrary.dirty && document.getElementById("autosave-status").dataset.state === "saved"')


def preset_saved(page):
    page.wait_for_function('window.journalWorkspace?.ready && !journalWorkspace.dirty && document.getElementById("journal-save-status").dataset.state === "saved"')


def login(page, url, username, register=False, display=None):
    page.goto(url, wait_until='networkidle')
    if register:
        page.locator('#register-tab').click()
        page.locator('#auth-display-name').fill(display or username)
    page.locator('#auth-username').fill(username)
    page.locator('#auth-password').fill(PASSWORD)
    page.locator('#auth-submit').click()
    ready(page)


def api(page, path):
    return page.evaluate('(path) => aiditorFetch(path).then(r=>r.json())', path)


def run():
    with tempfile.TemporaryDirectory(prefix='aiditor-browser-') as data, sync_playwright() as p:
        server=Server(data)
        try:
            url=server.start()
            browser=p.chromium.launch(headless=True)
            context=browser.new_context(viewport={'width':1440,'height':1050})
            page=context.new_page()
            errors=[]
            page.on('pageerror',lambda error:errors.append(str(error)))
            page.on('dialog',lambda dialog:dialog.accept())
            page.goto(url,wait_until='networkidle')
            page.screenshot(path=str(QA/'01-login.png'),full_page=True)
            assert 'MIT' in page.locator('footer').inner_text()
            login(page,url,'journal_alpha',True,'Bilim ve Toplum Dergisi')
            expect(page.locator('.template-option')).to_have_count(4)
            for template in ['classic','contemporary','centered','minimal']:
                page.locator(f'[data-template-id="{template}"]').click()
                preset_saved(page)
                assert api(page,'/api/journal')['settings']['template_id']==template
            page.locator('[data-template-id="contemporary"]').click()
            page.locator('#js-name-en').fill('Science and Society Review')
            page.locator('#js-url').fill('https://example.org/science')
            page.locator('#js-footer-text').fill('Dergiye özel yayın notu · CC BY 4.0')
            page.locator('#js-first-page-fit').select_option('compact')
            logo=Path(data)/'dergi-logom.png'
            Image.new('RGB',(180,90),(22,94,95)).save(logo)
            page.locator('#logo-inp').set_input_files(logo)
            preset_saved(page)
            journal=api(page,'/api/journal')
            assert journal['settings']['footer_text'].startswith('Dergiye özel')
            assert journal['assets']['logo']['name']=='dergi-logom.png'
            assert len(journal['assets']['logo']['data'])>100
            # Imported preset must be portable, including actual logo bytes.
            with page.expect_download() as download:
                page.locator('#export-preset').click()
            preset=Path(data)/'journal-preset.json';download.value.save_as(preset)
            portable=json.loads(preset.read_text())
            assert portable['format']=='aiditor-journal-preset'
            assert portable['assets']['logo']==journal['assets']['logo']
            page.screenshot(path=str(QA/'02-journal-settings.png'),full_page=True)
            # A save error remains visible and preserves edits.
            def fail_settings(route):
                if route.request.method=='PUT':route.fulfill(status=503,json={'ok':False,'error':'Test: disk unavailable'})
                else:route.continue_()
            page.route('**/api/journal',fail_settings)
            page.locator('#js-footer-text').fill('Korunması gereken son dipnot')
            expect(page.locator('#journal-save-status')).to_have_attribute('data-state','error')
            assert page.evaluate('journalWorkspace.prepareToClose()') is False
            page.unroute('**/api/journal',fail_settings)
            page.locator('#save-journal').click();preset_saved(page)
            # Reload restores preset, logo and enum values.
            page.reload(wait_until='networkidle');ready(page)
            page.locator('#journal-tab').click()
            expect(page.locator('#js-footer-text')).to_have_value('Korunması gereken son dipnot')
            expect(page.locator('#logo-label')).to_have_text('dergi-logom.png')
            # Restore exported preferences without affecting article drafts.
            page.locator('#import-preset').set_input_files(preset)
            expect(page.locator('#js-footer-text')).to_have_value('Dergiye özel yayın notu · CC BY 4.0')
            preset_saved(page)
            page.locator('#articles-tab').click()
            page.locator('#new-article').click()
            expect(page.locator('#c-first-page-fit')).to_have_value('compact')
            page.locator('#c-tr-title').fill('Yarım kalan çalışma <img src=x onerror=alert(1)>')
            page.locator('#c-en-title').fill('Open editorial workflows')
            page.locator('#authors-list [data-f="name"]').first.fill('İpek Çelik')
            page.locator('#authors-list [data-f="orcid"]').first.fill('0000-')
            page.locator('#c-doi').fill('10.')
            saved(page)
            first_id=page.evaluate('articleLibrary.currentId')
            assert page.locator('#article-list img').count()==0
            # Rich HTML paste preserves paragraph/table placement and merged cells.
            page.evaluate('''() => {
                const el=document.querySelector('#sections-list [data-f="content"]');
                const dt=new DataTransfer();
                dt.setData('text/html','<p>Tablodan önceki metin.</p><table><tr><th colspan="2">Birleşik başlık</th></tr><tr><td><b>Değer</b></td><td>42</td></tr></table><p>Tablodan sonraki metin.</p>');
                dt.setData('text/plain','Tablodan önceki metin.');
                el.dispatchEvent(new ClipboardEvent('paste',{bubbles:true,clipboardData:dt}));
            }''')
            saved(page)
            table=page.evaluate('collectFTs().find(ft=>ft.type==="table")')
            assert table['tbl_model']['rows'][0][0]['colspan']==2
            assert table['section_id'] is not None
            figure_id=page.evaluate('String(addFT({type:"figure"}))')
            page.locator(f'#finp-{figure_id}').set_input_files(logo)
            saved(page)
            # Last keystrokes flush before switching articles.
            page.locator('#refs').fill('Kaybolmaması gereken son kaynak')
            page.locator('#new-article').click()
            expect(page.locator('#c-tr-title')).to_have_value('')
            page.locator('#c-tr-title').fill('İkinci makale');saved(page)
            second_id=page.evaluate('articleLibrary.currentId')
            assert second_id!=first_id
            page.locator(f'[data-article-id="{first_id}"]').click()
            expect(page.locator('#refs')).to_have_value('Kaybolmaması gereken son kaynak')
            expect(page.locator('#c-doi')).to_have_value('10.')
            assert page.evaluate('Object.keys(figFiles).length')==1
            # Two-window stale writes become recoverable copies.
            other=context.new_page();other.on('dialog',lambda dialog:dialog.accept())
            other.goto(url,wait_until='networkidle');ready(other)
            if other.evaluate('articleLibrary.currentId')!=first_id:
                other.locator(f'[data-article-id="{first_id}"]').click()
            page.locator('#refs').fill('Ana penceredeki güncel kaynak');saved(page)
            other.locator('#refs').fill('Diğer penceredeki ayrı değişiklik')
            expect(other.locator('#save-conflict-copy')).to_be_visible()
            other.locator('#save-conflict-copy').click();saved(other)
            assert other.evaluate('articleLibrary.currentId')!=first_id
            assert api(page,'/api/articles/'+first_id)['article']['project']['data']['references']=='Ana penceredeki güncel kaynak'
            other.close()
            # Disk failure cannot lose current draft by switching away.
            def fail_article(route):
                if route.request.method=='PUT':route.fulfill(status=503,json={'ok':False,'error':'Test: unavailable'})
                else:route.continue_()
            page.route('**/api/articles/*',fail_article)
            page.locator('#refs').fill('Kesintide korunan kaynak')
            expect(page.locator('#autosave-status')).to_have_attribute('data-state','error')
            page.locator('#new-article').click()
            expect(page.locator('#refs')).to_have_value('Kesintide korunan kaynak')
            page.unroute('**/api/articles/*',fail_article)
            page.locator('#autosave-retry').click();saved(page)
            # The exported JSON retains incomplete DOI/ORCID and rich tables.
            with page.expect_download() as download:
                page.locator('button[onclick="saveProject()"]').click()
            project=Path(data)/'draft.json';download.value.save_as(project)
            exported=json.loads(project.read_text())
            assert exported['format']=='aiditor-project'
            assert exported['data']['cover']['doi']=='10.'
            assert len(exported['figures'])==1
            page.locator('#c-doi').fill('10.1234/example.2026')
            page.locator('#authors-list [data-f="orcid"]').first.fill('')
            saved(page)
            page.locator('#btn-gen').click()
            page.wait_for_selector('#result-panel',state='visible')
            link=page.locator('#dl-link')
            response=context.request.get(url+link.get_attribute('href'))
            assert response.status==200
            with zipfile.ZipFile(io.BytesIO(response.body())) as archive:
                assert archive.testzip() is None
                tex=archive.read('main.tex').decode()
                assert 'Science and Society Review' in tex
                assert 'journal_logo.png' in archive.namelist()
                assert 'JGTTR.png' not in archive.namelist()
            page.screenshot(path=str(QA/'03-article-editor.png'),full_page=True)
            # DOCX import uses source ordering and stable table links.
            doc=Document();doc.add_paragraph('Belge Başlığı',style='Title');doc.add_heading('Giriş',level=1)
            doc.add_paragraph('Wordden gelen ilk paragraf.');t=doc.add_table(rows=2,cols=2)
            t.cell(0,0).text='Başlık';t.cell(0,1).text='Değer';t.cell(1,0).text='Veri';t.cell(1,1).text='7'
            doc.add_paragraph('Tablo sonrasındaki paragraf.');docfile=Path(data)/'article.docx';doc.save(docfile)
            page.locator('#docx-file-input').set_input_files(docfile)
            page.wait_for_function('!docxImporting && collectSections().some(s=>s.content.includes("Wordden gelen"))')
            saved(page)
            assert page.evaluate('collectFTs().some(ft=>ft.type==="table" && ft.section_id!=null)')
            # Restore a portable project after preserving the imported article.
            page.locator('#project-file').set_input_files(project)
            page.wait_for_function('document.getElementById("c-doi").value==="10."')
            saved(page)
            # A stale tab must never read/write into the next account.
            stale=context.new_page();stale.on('dialog',lambda dialog:dialog.accept());stale.goto(url,wait_until='networkidle');ready(stale)
            page.locator('#logout-button').click();expect(page.locator('#auth-screen')).to_be_visible()
            login(page,url,'journal_beta',True,'Başka Dergi')
            assert api(page,'/api/articles')['articles']==[]
            assert api(page,'/api/journal')['assets']=={}
            assert page.evaluate('(id)=>aiditorFetch("/api/articles/"+id).then(r=>r.status)',first_id)==404
            stale.locator('#refs').fill('Eski hesapta kalmalı')
            expect(stale.locator('#session-warning')).to_be_visible()
            assert api(page,'/api/articles')['articles']==[]
            stale.close()
            page.locator('#journal-tab').click()
            page.locator('#import-preset').set_input_files(preset)
            expect(page.locator('#logo-label')).to_have_text('dergi-logom.png')
            preset_saved(page)
            assert api(page,'/api/journal')['assets']['logo']['name']=='dergi-logom.png'
            # Invalid preset imports preserve both the working fields and disk record.
            invalid=Path(data)/'invalid-preset.json'
            invalid.write_text(json.dumps({'format':'aiditor-journal-preset','version':1,'settings':{'journal_name_tr':'REPLACEMENT','accent_color':'invalid'},'assets':{}}))
            with page.expect_response('**/validate_journal') as preflight:
                page.locator('#import-preset').set_input_files(invalid)
            assert preflight.value.status==400
            expect(page.locator('#js-name-tr')).to_have_value('Bilim ve Toplum Dergisi')
            assert api(page,'/api/journal')['settings']['journal_name_tr']=='Bilim ve Toplum Dergisi'
            # Mobile layout has no page-wide horizontal overflow.
            page.set_viewport_size({'width':390,'height':844})
            page.screenshot(path=str(QA/'04-mobile-settings.png'),full_page=True)
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth+1')
            page.locator('#logout-button').click();expect(page.locator('#auth-screen')).to_be_visible()
            page.screenshot(path=str(QA/'05-mobile-login.png'),full_page=True)
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth+1')
            context.close()
            # Restart the service on another port; no browser session or account data in test repo.
            server.stop();url=server.start()
            context=browser.new_context();page=context.new_page();page.on('dialog',lambda dialog:dialog.accept())
            login(page,url,'journal_alpha')
            assert api(page,'/api/journal')['settings']['journal_name_en']=='Science and Society Review'
            assert api(page,'/api/journal')['assets']['logo']['name']=='dergi-logom.png'
            assert len(api(page,'/api/articles')['articles'])>=3
            assert not errors,errors
            browser.close()
            print('PASS: accounts, four presets, logo, recovery, autosave, conflicts, rich paste, DOCX, ZIP, JSON, isolation, restart, mobile')
        finally:server.stop()


if __name__=='__main__':
    run()
