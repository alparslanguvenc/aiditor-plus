"""Select the neutral bilingual preset and export both editable and Overleaf formats."""
import io
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from playwright.sync_api import expect, sync_playwright

from browser_workflow import Server, login, preset_saved, ready


def run():
    with tempfile.TemporaryDirectory(prefix='aiditor-panel-browser-') as data, sync_playwright() as playwright:
        server = Server(data)
        browser = None
        try:
            url = server.start()
            browser = playwright.chromium.launch(headless=True)
            page = browser.new_page(viewport={'width': 1280, 'height': 900})
            errors = []
            page.on('pageerror', lambda error: errors.append(str(error)))
            login(page, url, 'panel_test', True, 'Örnek Dergi')
            page.locator('#journal-tab').click()
            page.locator('#js-name-tr').fill('Kendi Dergim')
            page.locator('#js-issn-online').fill('1234-567X')
            page.locator('#js-url').fill('https://example.org/my-journal')
            preset_saved(page)
            page.locator('[data-template-id="bilingual_panel"]').click()
            preset_saved(page)
            expect(page.locator('#js-name-tr')).to_have_value('Kendi Dergim')
            expect(page.locator('#js-issn-online')).to_have_value('1234-567X')
            expect(page.locator('#js-url')).to_have_value('https://example.org/my-journal')
            expect(page.locator('#bilingual-panel-layout-note')).to_be_visible()
            assert page.locator('[data-template-id="bilingual_panel"] img').evaluate(
                '(img) => img.complete && img.naturalWidth > 0')
            page.reload(); ready(page); page.locator('#journal-tab').click()
            expect(page.locator('[data-template-id="bilingual_panel"]')).to_have_attribute('aria-pressed', 'true')
            page.locator('#articles-tab').click()
            fields = {'c-tr-title': 'Örnek Türkçe Başlık', 'c-en-title': 'Example English Title',
                      'a-tr-abs': 'Türkçe özet alanı.', 'a-en-abs': 'English abstract panel.'}
            page.evaluate('''fields => { for (const [id, value] of Object.entries(fields)) {
                const element = document.getElementById(id); element.value = value;
                element.dispatchEvent(new Event('input', {bubbles: true})); }}''', fields)
            page.locator('#btn-docx').click()
            expect(page.locator('#result-title')).to_have_text('Tamamlandı — Word belgesi hazır')
            with page.expect_download() as download:
                page.locator('#dl-docx').click()
            path = Path(data) / 'panel.docx'; download.value.save_as(path)
            doc = Document(io.BytesIO(path.read_bytes()))
            assert len(doc.tables) >= 3
            assert 'Kendi Dergim' in doc.tables[0].cell(0, 1).text
            assert fields['a-tr-abs'] in doc.tables[1].cell(0, 1).text
            assert fields['a-en-abs'] in doc.tables[2].cell(0, 1).text
            page.locator('#btn-gen').click()
            expect(page.locator('#dl-link')).to_be_visible()
            with page.expect_download() as download:
                page.locator('#dl-link').click()
            path = Path(data) / 'panel.zip'; download.value.save_as(path)
            with zipfile.ZipFile(path) as archive:
                assert archive.testzip() is None
                tex = archive.read('main.tex').decode()
                assert '% Template: bilingual_panel;' in tex
                assert 'Kendi Dergim' in tex
                assert 'Gastroia' not in tex
            assert not errors, errors
            print('PASS: neutral panel preview, identity persistence, DOCX panels, Overleaf ZIP')
        finally:
            if browser:
                browser.close()
            server.stop()


if __name__ == '__main__':
    run()
