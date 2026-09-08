"""Behavioural coverage for local accounts, durable settings and isolated drafts."""
from contextlib import closing
import base64
import copy
import io
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import uuid
import zipfile
import unittest
from docx import Document
os.environ.setdefault('AIDITOR_DATA_DIR', tempfile.mkdtemp(prefix='aiditor-test-bootstrap-'))
import app as application
from account_store import AccountStore, RevisionConflict
from journal_templates import default_settings
PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=')
ASSET = {'name': 'journal.png', 'data': 'data:image/png;base64,' + base64.b64encode(PNG).decode()}

def make_setup(tmp_path):
    application.app.config.update(TESTING=True, AIDITOR_DATA_DIR=tmp_path, SECRET_KEY=AccountStore(tmp_path).session_secret(), LEGACY_PROFILES_DIR=tmp_path / 'legacy')
    application._zip_store.clear()
    application._docx_import_store.clear()
    return (application.app, tmp_path)

def register(client, username='journal_a'):
    response = client.post('/api/auth/register', json={'username': username, 'password': 'correct-password', 'display_name': username.upper()}, headers={'X-Aiditor-Request': '1'})
    assert response.status_code == 201, response.json
    return response.json['user']

def headers(user):
    return {'X-Aiditor-Request': '1', 'X-Aiditor-Account': user['id']}

def project(title='Draft title'):
    return {'format': 'aiditor-project', 'version': 1, 'data': {'cover': {'tr_title': title}, 'abstract': {}, 'extra': {}, 'authors': [], 'sections': [], 'figtables': [], 'references': ''}, 'figures': {}}

class TestAccounts(unittest.TestCase):

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='aiditor-accounts-test-')
        self.addCleanup(self.temporary.cleanup)
        self.setup = make_setup(Path(self.temporary.name))

    def test_account_password_hash_session_cookie_and_restart(self):
        setup = self.setup
        app, directory = setup
        client = app.test_client()
        user = register(client)
        cookie = client.get_cookie('aiditor_plus_session')
        assert cookie.http_only and cookie.same_site == 'Strict'
        with closing(sqlite3.connect(directory / 'aiditor.sqlite3')) as db:
            stored = db.execute('SELECT password_hash FROM accounts').fetchone()[0]
        assert stored.startswith('scrypt:') and stored != 'correct-password'
        store = AccountStore(directory)
        assert store.authenticate('journal_a', 'correct-password') == user
        assert store.authenticate('journal_a', 'wrong-password') is None
        assert store.session_secret() == AccountStore(directory).session_secret()
        assert client.get('/api/auth/session').json['user'] == user

    def test_credentials_and_duplicate_account(self):
        setup = self.setup
        app, _ = setup
        client = app.test_client()
        for username, password in [('x', 'long-password'), ('has space', 'long-password'), ('valid', 'short')]:
            response = client.post('/api/auth/register', json={'username': username, 'password': password}, headers={'X-Aiditor-Request': '1'})
            assert response.status_code == 400
        register(client)
        response = client.post('/api/auth/register', json={'username': 'JOURNAL_A', 'password': 'long-password'}, headers={'X-Aiditor-Request': '1'})
        assert response.status_code == 409
        response = client.post('/api/auth/login', json={'username': 'journal_a', 'password': 'wrong-password'}, headers={'X-Aiditor-Request': '1'})
        assert response.status_code == 401

    def test_settings_assets_persist_and_are_isolated(self):
        setup = self.setup
        app, directory = setup
        a, b = (app.test_client(), app.test_client())
        ua, ub = (register(a), register(b, 'journal_b'))
        settings = default_settings()
        settings.update(journal_name_tr='Kalıcı Dergi', footer_text='Yayın kurulu dipnotu', template_id='centered')
        saved = a.put('/api/journal', json={'settings': settings, 'assets': {'logo': ASSET}, 'base_revision': 0}, headers=headers(ua))
        assert saved.status_code == 200, saved.json
        assert saved.json['revision'] == 1
        restarted = AccountStore(directory).journal(ua['id'])
        assert restarted['settings']['journal_name_tr'] == 'Kalıcı Dergi'
        assert restarted['assets']['logo'] == ASSET
        other = b.get('/api/journal', headers=headers(ub)).json
        assert other['settings']['journal_name_tr'] == 'JOURNAL_B'
        assert other['assets'] == {}
        assert 'JGTTR' not in json.dumps(other)

    def test_journal_revision_conflict_keeps_latest(self):
        setup = self.setup
        app, _ = setup
        client = app.test_client()
        user = register(client)
        settings = default_settings()
        settings['journal_name_tr'] = 'New title'
        assert client.put('/api/journal', json={'settings': settings, 'base_revision': 0}, headers=headers(user)).status_code == 200
        settings['journal_name_tr'] = 'Stale title'
        response = client.put('/api/journal', json={'settings': settings, 'base_revision': 0}, headers=headers(user))
        assert response.status_code == 409 and response.json['code'] == 'revision_conflict'
        assert client.get('/api/journal', headers=headers(user)).json['settings']['journal_name_tr'] == 'New title'

    def test_stale_tab_cannot_write_or_read_other_journal(self):
        setup = self.setup
        app, _ = setup
        client = app.test_client()
        old = register(client)
        new = register(client, 'journal_b')
        for response in [client.get('/api/journal', headers=headers(old)), client.get('/api/articles', headers=headers(old)), client.put('/api/journal', json={'settings': default_settings(), 'base_revision': 0}, headers=headers(old)), client.post('/api/auth/logout', json={}, headers=headers(old))]:
            assert response.status_code == 409 and response.json['code'] == 'account_changed'
        assert client.get('/api/auth/session').json['user'] == new

    def test_private_routes_require_login(self):
        setup = self.setup
        for route, method in [('/api/journal', 'get'), ('/api/articles', 'get'), ('/process_form', 'post'), ('/import_docx', 'post')]:
            with self.subTest():
                app, _ = setup
                response = getattr(app.test_client(), method)(route, headers={'X-Aiditor-Request': '1'})
                assert response.status_code == 401

    def test_origin_host_and_request_header_protection(self):
        setup = self.setup
        app, _ = setup
        client = app.test_client()
        user = register(client)
        payload = {'settings': default_settings(), 'base_revision': 0}
        assert client.put('/api/journal', json=payload, headers={'X-Aiditor-Account': user['id']}).status_code == 403
        bad = dict(headers(user), Origin='https://evil.example')
        assert client.put('/api/journal', json=payload, headers=bad).status_code == 403
        assert client.get('/health', base_url='http://evil.example').status_code == 403
        assert client.put('/api/journal', json=payload, headers=dict(headers(user), Origin='http://localhost')).status_code == 200

    def test_article_restart_conflict_retry_and_delete(self):
        setup = self.setup
        app, directory = setup
        client = app.test_client()
        user = register(client)
        identifier = str(uuid.uuid4())
        payload = {'project': project(), 'base_revision': 0}
        response = client.put('/api/articles/' + identifier, json=payload, headers=headers(user))
        assert response.status_code == 200 and response.json['article']['revision'] == 1
        retry = client.put('/api/articles/' + identifier, json=payload, headers=headers(user))
        assert retry.json['article']['revision'] == 1
        assert AccountStore(directory).article(user['id'], identifier)['project'] == payload['project']
        payload['project']['data']['cover']['tr_title'] = 'Stale changed draft'
        assert client.put('/api/articles/' + identifier, json=payload, headers=headers(user)).status_code == 409
        assert client.delete('/api/articles/' + identifier, json={'base_revision': 0}, headers=headers(user)).status_code == 409
        assert client.delete('/api/articles/' + identifier, json={'base_revision': 1}, headers=headers(user)).status_code == 200
        assert client.get('/api/articles/' + identifier, headers=headers(user)).status_code == 404

    def test_article_owner_isolation_even_same_uuid(self):
        setup = self.setup
        app, _ = setup
        a, b = (app.test_client(), app.test_client())
        ua, ub = (register(a), register(b, 'journal_b'))
        identifier = str(uuid.uuid4())
        route = '/api/articles/' + identifier
        assert a.put(route, json={'project': project('Secret A'), 'base_revision': 0}, headers=headers(ua)).status_code == 200
        assert b.get(route, headers=headers(ub)).status_code == 404
        assert b.delete(route, json={'base_revision': 1}, headers=headers(ub)).status_code == 404
        assert b.get('/api/articles', headers=headers(ub)).json['articles'] == []
        assert b.put(route, json={'project': project('Private B'), 'base_revision': 0}, headers=headers(ub)).status_code == 200
        assert a.get(route, headers=headers(ua)).json['article']['title'] == 'Secret A'

    def test_bad_journal_assets_rejected_without_persistence(self):
        setup = self.setup
        for asset in [{'name': '../outside.png', 'data': ASSET['data']}, {'name': 'fake.jpg', 'data': ASSET['data']}, {'name': 'fake.png', 'data': 'data:image/png;base64,YmFk'}, {'name': 'logo.svg', 'data': 'data:image/svg+xml;base64,PHN2Zz48L3N2Zz4='}]:
            with self.subTest():
                app, _ = setup
                client = app.test_client()
                user = register(client, 'journal_' + uuid.uuid4().hex[:8])
                response = client.put('/api/journal', json={'settings': default_settings(), 'base_revision': 0, 'assets': {'logo': asset}}, headers=headers(user))
                assert response.status_code == 400
                assert client.get('/api/journal', headers=headers(user)).json['assets'] == {}

    def test_invalid_shapes_and_project_figure_validation(self):
        setup = self.setup
        app, _ = setup
        client = app.test_client()
        user = register(client)
        assert client.put('/api/journal', json=[], headers=headers(user)).status_code == 400
        assert client.put('/api/journal', json={'settings': {}, 'base_revision': True}, headers=headers(user)).status_code == 400
        malformed = project()
        malformed['data']['authors'] = 'not a list'
        route = '/api/articles/' + str(uuid.uuid4())
        assert client.put(route, json={'project': malformed, 'base_revision': 0}, headers=headers(user)).status_code == 400
        missing = project()
        missing['data']['figtables'] = [{'type': 'figure', 'number': '1', 'file_key': 'absent'}]
        assert client.put(route, json={'project': missing, 'base_revision': 0}, headers=headers(user)).status_code == 400
        missing['data']['figtables'][0]['file_missing'] = True
        assert client.put(route, json={'project': missing, 'base_revision': 0}, headers=headers(user)).status_code == 200
        assert client.get('/api/articles/not-a-uuid', headers=headers(user)).status_code == 400

    def test_jgttr_project_compatibility_without_rewriting_draft(self):
        setup = self.setup
        app, _ = setup
        client = app.test_client()
        user = register(client)
        draft = project()
        draft['format'] = 'jgttr-project'
        draft['data']['authors'] = [{'name': 'A', 'orcid': 'unfinished'}]
        identifier = str(uuid.uuid4())
        response = client.put('/api/articles/' + identifier, json={'project': draft, 'base_revision': 0}, headers=headers(user))
        assert response.status_code == 200
        assert client.get('/api/articles/' + identifier, headers=headers(user)).json['article']['project'] == draft

    def test_saved_assets_generation_and_owner_download(self):
        setup = self.setup
        app, _ = setup
        a, b = (app.test_client(), app.test_client())
        ua, ub = (register(a), register(b, 'journal_b'))
        settings = default_settings()
        settings['journal_name_tr'] = 'Örnek Araştırma Dergisi'
        a.put('/api/journal', json={'settings': settings, 'base_revision': 0, 'assets': {'logo': ASSET, 'license': ASSET}}, headers=headers(ua))
        response = a.post('/process_form', data={'data': json.dumps(project()['data'])}, headers=headers(ua))
        assert response.status_code == 200, response.json
        route = '/download/' + response.json['key']
        assert b.get(route).status_code == 404
        downloaded = a.get(route)
        assert downloaded.status_code == 200
        with zipfile.ZipFile(io.BytesIO(downloaded.data)) as archive:
            assert archive.testzip() is None
            assert archive.read('journal_logo.png') == PNG
            assert archive.read('journal_license.png') == PNG
            assert 'Örnek Araştırma Dergisi' in archive.read('main.tex').decode()
            assert 'JGTTR.png' not in archive.namelist()
            assert b'MIT' in archive.read('README_Overleaf.txt')

    def test_validate_form_bad_input_and_file_paths(self):
        setup = self.setup
        app, _ = setup
        client = app.test_client()
        user = register(client)
        response = client.post('/validate_form', data={'data': '{invalid'}, headers=headers(user))
        assert response.status_code == 400
        data = project()['data']
        data['cover']['doi'] = 'bad DOI'
        assert client.post('/validate_form', data={'data': json.dumps(data)}, headers=headers(user)).status_code == 400
        assert client.post('/validate_form', data={'data': json.dumps(data), 'draft': '1'}, headers=headers(user)).status_code == 200
        assert client.post('/process_form', data={'data': json.dumps(project()['data']), 'fig_../bad': (io.BytesIO(PNG), 'x.png')}, headers=headers(user)).status_code == 400

    def test_docx_import_and_image_owner_isolation(self):
        setup = self.setup
        app, _ = setup
        a, b = (app.test_client(), app.test_client())
        ua, ub = (register(a), register(b, 'journal_b'))
        doc = Document()
        doc.add_heading('Bir araştırma', 0)
        doc.add_paragraph('Abstract')
        doc.add_paragraph('This is the research abstract.')
        doc.add_picture(io.BytesIO(PNG))
        stream = io.BytesIO()
        doc.save(stream)
        response = a.post('/import_docx', data={'article': (io.BytesIO(stream.getvalue()), 'article.docx')}, headers=headers(ua))
        assert response.status_code == 200, response.json
        key = response.json['import_key']
        assert a.get(f'/import_docx/{key}/image/0').status_code == 200
        assert b.get(f'/import_docx/{key}/image/0').status_code == 404
        assert a.post('/import_docx', data={'article': (io.BytesIO(b'bad archive'), 'article.docx')}, headers=headers(ua)).status_code == 400

    def test_legacy_migration_is_explicit_safe_and_non_destructive(self):
        setup = self.setup
        app, directory = setup
        client = app.test_client()
        user = register(client)
        legacy = directory / 'legacy'
        legacy.mkdir()
        original = json.dumps({'journal_name_tr': 'Legacy title', 'logo_filename': 'legacy.png', 'logo_stem': 'Old Journal logo'})
        (legacy / 'Old Journal.json').write_text(original)
        (legacy / 'legacy.png').write_bytes(PNG)
        (legacy / 'escape.json').symlink_to(directory / 'aiditor.sqlite3')
        profiles = client.get('/api/legacy-profiles', headers=headers(user)).json['profiles']
        assert profiles == ['Old Journal']
        response = client.get('/api/legacy-profiles/Old%20Journal', headers=headers(user))
        assert response.status_code == 200
        assert response.json['assets']['logo']['data'] == ASSET['data']
        assert client.get('/api/journal', headers=headers(user)).json['settings']['journal_name_tr'] == 'JOURNAL_A'
        payload = {'settings': response.json['settings'], 'assets': response.json['assets'], 'base_revision': 0}
        assert client.put('/api/journal', json=payload, headers=headers(user)).status_code == 200
        assert (legacy / 'Old Journal.json').read_text() == original
        assert client.delete('/delete_profile/Old%20Journal', headers=headers(user)).status_code == 410
        (legacy / 'danger.json').write_text(json.dumps({'logo_filename': '../private.png'}))
        assert client.get('/api/legacy-profiles/danger', headers=headers(user)).status_code == 400
        assert client.get('/api/legacy-profiles/escape', headers=headers(user)).status_code == 404
    def test_pdf_assets_parse_reject_spoof_and_encryption(self):
        from pypdf import PdfWriter
        app, _ = self.setup
        client = app.test_client()
        user = register(client)
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=100)
        stream = io.BytesIO()
        writer.write(stream)
        pdf = stream.getvalue()
        asset = {'name': 'logo.pdf', 'data': 'data:application/pdf;base64,' + base64.b64encode(pdf).decode()}
        response = client.put('/api/journal', json={'settings': default_settings(), 'base_revision': 0, 'assets': {'logo': asset}}, headers=headers(user))
        assert response.status_code == 200, response.json
        generated = client.post('/process_form', data={'data': json.dumps(project()['data'])}, headers=headers(user))
        assert generated.status_code == 200
        downloaded = client.get('/download/' + generated.json['key'])
        with zipfile.ZipFile(io.BytesIO(downloaded.data)) as archive:
            assert archive.read('journal_logo.pdf') == pdf
        spoof = {'name': 'spoof.pdf', 'data': 'data:application/pdf;base64,' + base64.b64encode(b'%PDF-1.7 fake document %%EOF').decode()}
        assert client.put('/api/journal', json={'settings': default_settings(), 'base_revision': 1, 'assets': {'logo': spoof}}, headers=headers(user)).status_code == 400
        assert client.post('/process_form', data={'data': json.dumps(project()['data']), 'fig_test': (io.BytesIO(b'%PDF-1.7 fake document %%EOF'), 'spoof.pdf')}, headers=headers(user)).status_code == 400
        writer.encrypt('secret')
        encrypted = io.BytesIO()
        writer.write(encrypted)
        asset['data'] = 'data:application/pdf;base64,' + base64.b64encode(encrypted.getvalue()).decode()
        assert client.put('/api/journal', json={'settings': default_settings(), 'base_revision': 1, 'assets': {'logo': asset}}, headers=headers(user)).status_code == 400

    def test_license_is_public(self):
        app, _ = self.setup
        response = app.test_client().get('/license')
        self.addCleanup(response.close)
        assert response.status_code == 200
        assert b'MIT License' in response.data

    def test_journal_import_preflight_does_not_mutate_saved_state(self):
        app, _ = self.setup
        client = app.test_client()
        user = register(client)
        before = client.get('/api/journal', headers=headers(user)).json
        incoming = dict(default_settings(), journal_name_tr='Imported journal')
        response = client.post('/validate_journal', json={'settings': incoming, 'assets': {'logo': ASSET}}, headers=headers(user))
        assert response.status_code == 200
        assert response.json['settings']['journal_name_tr'] == 'Imported journal'
        incoming['accent_color'] = 'invalid'
        assert client.post('/validate_journal', json={'settings': incoming, 'assets': {'logo': ASSET}}, headers=headers(user)).status_code == 400
        assert client.get('/api/journal', headers=headers(user)).json == before

if __name__ == '__main__':
    unittest.main()
