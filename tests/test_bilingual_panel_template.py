"""The supplied DOCX contributes layout only; journal identity comes from settings."""
import io
import unittest

from docx import Document

from docx_export import generate_docx_from_form
from formatter import generate_latex_from_form
from journal_templates import TEMPLATES, normalize_settings
from page_furniture import running_slots
from test_accounts import PNG
from test_formatter_templates import sample_article


SETTINGS = next(item['settings'] for item in TEMPLATES if item['id'] == 'bilingual_panel')


class BilingualPanelTests(unittest.TestCase):
    def test_editable_cover_has_two_panels_and_user_identity(self):
        data = sample_article()
        preferences = {**SETTINGS, 'journal_name_tr': 'Örnek Dergi',
                       'journal_name_en': 'Example Journal',
                       'journal_url': 'https://example.org/journal',
                       'issn_online': '1234-567X', 'accent_color': '#336699'}
        doc = Document(io.BytesIO(generate_docx_from_form(data, {}, preferences,
                                      {'logo': ('own-logo.png', PNG)})))
        cover, body = doc.sections
        self.assertAlmostEqual(cover.left_margin.cm, 2, places=2)
        self.assertAlmostEqual(body.right_margin.cm, 2, places=2)
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertEqual(len(doc.tables), 4)  # masthead, two panels, article table
        self.assertIn('Örnek Dergi', doc.tables[0].cell(0, 1).text)
        self.assertIn('https://example.org/journal', doc.tables[0].cell(0, 1).text)
        self.assertIn(data['abstract']['tr_abs'], doc.tables[1].cell(0, 1).text)
        self.assertIn(data['abstract']['en_abs'], doc.tables[2].cell(0, 1).text)
        self.assertIn('D6E0EB', doc.tables[0]._tbl.xml.upper())
        self.assertIn('e-ISSN: 1234-567X', cover.first_page_header.tables[0].cell(0, 0).text)
        self.assertIn('Örnek Dergi', body.even_page_header.tables[0].cell(0, 0).text)
        self.assertIn('PAGE', cover.first_page_footer._element.xml)
        self.assertNotIn('Gastroia', doc.element.xml)

    def test_latex_has_same_panels_and_custom_running_furniture(self):
        data = sample_article()
        preferences = {**SETTINGS, 'journal_name_tr': 'Örnek Dergi',
                       'issn_online': '1234-567X', 'footer_text': 'Özel yayın notu'}
        tex = generate_latex_from_form(data, {}, preferences)
        self.assertIn('% Template: bilingual_panel;', tex)
        self.assertIn(r'\colorbox{JGTTRbrown!30}', tex)
        self.assertEqual(tex.count(r'\colorbox{JGTTRgray!35}'), 2)
        self.assertIn(r'\JGTTRturkishabstract', tex)
        self.assertIn(r'\JGTTRenglishabstract', tex)
        self.assertIn('left=2cm,right=2cm', tex)
        self.assertIn('e-ISSN: 1234-567X', tex)
        self.assertIn('Özel yayın notu', tex)
        self.assertNotIn('Gastroia', tex)
        self.assertNotIn('2602-4144', tex)

    def test_english_only_omits_turkish_panel(self):
        data = sample_article()
        settings = {**SETTINGS, 'english_only': True}
        doc = Document(io.BytesIO(generate_docx_from_form(data, {}, settings)))
        text = '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        self.assertNotIn(data['abstract']['tr_abs'], text)
        self.assertIn(data['abstract']['en_abs'], text)
        tex = generate_latex_from_form(data, {}, settings)
        layout = tex[tex.index('% Template: bilingual_panel;'):]
        self.assertNotIn(r'\JGTTRturkishabstract', layout)
        self.assertEqual(layout.count(r'\JGTTRenglishabstract'), 1)

    def test_missing_issn_does_not_print_empty_label(self):
        self.assertEqual(running_slots(normalize_settings(SETTINGS), sample_article(), 'header', 'first')[1],
                         '{cilt}({sayi}) · {yil}')


if __name__ == '__main__':
    unittest.main()
