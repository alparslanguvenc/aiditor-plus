import base64

import io

import unittest

from docx import Document

from docx.enum.style import WD_STYLE_TYPE

from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.oxml import OxmlElement

from docx.oxml.ns import qn

from docx.shared import Cm

from formatter import (
    _build_figtable_latex,
    _escape_with_breakable_urls,
    escape,
    extract_form_data_from_docx,
    generate_latex_from_form,
)

def minimal_form_data(**cover_overrides):
    cover = {
        'tr_title': 'Türkçe Başlık',
        'en_title': 'English Title',
        'year': '2026',
        'first_page_fit': 'auto',
    }
    cover.update(cover_overrides)
    return {
        'cover': cover,
        'authors': [],
        'abstract': {
            'tr_abs': 'Kısa Türkçe özet.', 'tr_kw': 'turizm; teknoloji',
            'en_abs': 'A short abstract.', 'en_kw': 'tourism; technology',
        },
        'sections': [],
        'figtables': [],
        'extra': {},
        'references': '',
    }

class RichTableTests(unittest.TestCase):
    def test_rich_table_preserves_spans_alignment_and_emphasis(self):
        table = {
            'type': 'table', 'number': '2', 'tr_cap': 'Örnek & Tablo',
            'tbl_model': {
                'header_rows': 1,
                'grid_borders': True,
                'column_widths': [0.35, 0.65],
                'rows': [
                    [{'text': 'Birleşik Başlık', 'colspan': 2, 'bold': True, 'underline': True,
                      'align': 'center', 'bgcolor': 'D9EAF7', 'textcolor': '833C0B'}],
                    [
                        {'text': 'A & B', 'rowspan': 2, 'italic': True, 'align': 'left'},
                        {'text': '10', 'align': 'right'},
                    ],
                    [{'text': '20', 'align': 'right'}],
                ],
            },
        }

        latex = _build_figtable_latex(table, {})

        self.assertIn(r'\multicolumn{2}', latex)
        self.assertNotIn(r'\multirow', latex)
        self.assertIn(r'\textit{A \& B}', latex)
        self.assertIn(r'\cellcolor[HTML]{D9EAF7}', latex)
        self.assertIn(r'\textcolor[HTML]{833C0B}', latex)
        self.assertIn(r'\uline{', latex)
        self.assertIn(r'\caption{Örnek \& Tablo}', latex)
        self.assertIn(r'>{\raggedleft\arraybackslash}', latex)
        self.assertIn(r'\setlength{\tabcolsep}{3pt}', latex)
        self.assertIn(r'\begin{adjustbox}{max width=\linewidth,center}', latex)
        self.assertIn(r'\begin{tabular}{|', latex)
        self.assertIn(r'\setlength{\JGTTRtablewidth}{\linewidth}', latex)
        self.assertIn(r'\JGTTRtablewidth', latex)
        self.assertIn(r'\hline', latex)

    def test_multiline_cells_wrap_inside_their_columns(self):
        table = {
            'type': 'table', 'number': '3', 'split_table': True,
            'tbl_model': {
                'header_rows': 1,
                'grid_borders': True,
                'column_widths': [0.5, 0.5],
                'rows': [
                    [{'text': 'Başlık A'}, {'text': 'Başlık B'}],
                    [{
                        'text': (
                            'Birinci Word satırı\nİkinci Word satırı, hücre genişliğinde '
                            'doğal olarak sözcük sözcük sarılabilecek kadar uzundur.'
                        ),
                        'underline': True,
                    }, {'text': 'Değer'}],
                ],
            },
        }

        latex = _build_figtable_latex(table, {})

        self.assertIn(r'Birinci Word satırı\newline İkinci Word satırı', latex)
        self.assertNotIn(r'\makecell', latex)
        self.assertIn(r'\uline{', latex)

    def test_unicode_subscript_digits_are_renderable(self):
        self.assertEqual(escape('H₂ ve H₁₀'), r'H\textsubscript{2} ve H\textsubscript{1}\textsubscript{0}')

    def test_reference_urls_remain_breakable(self):
        latex = _escape_with_breakable_urls(
            'Rapor: https://example.org/a_very_long-path/report.pdf, erişim.'
        )
        self.assertIn(r'https:/\allowbreak{}/\allowbreak{}example.\allowbreak{}org', latex)
        self.assertIn(r'a\_\allowbreak{}very\_\allowbreak{}long', latex)
        self.assertIn(', erişim.', latex)

    def test_legacy_pipe_table_still_generates(self):
        table = {
            'type': 'table', 'number': '1',
            'tbl_data': 'Başlık 1 | Başlık 2\nDeğer | 25%',
        }
        latex = _build_figtable_latex(table, {})
        self.assertIn(r'\textbf{Başlık 1}', latex)
        self.assertIn(r'25\%', latex)

class DocxImportTests(unittest.TestCase):
    def build_article_docx(self):
        document = Document()
        document.add_paragraph('Sürdürülebilir Turizmde Dijital Dönüşüm', style='Title')
        document.add_paragraph('Digital Transformation in Sustainable Tourism', style='Subtitle')

        document.add_heading('Özet', level=1)
        document.add_paragraph('Bu çalışma Word içe aktarma akışını sınar.')
        document.add_paragraph('Anahtar Kelimeler: turizm; dijitalleşme')
        document.add_heading('Abstract', level=1)
        document.add_paragraph('This study tests the Word import workflow.')
        document.add_paragraph('Keywords: tourism; digitalisation')

        document.add_heading('Giriş / Introduction', level=1)
        document.add_paragraph('Araştırmanın başlangıç paragrafıdır.')
        document.add_paragraph('Tablo 1: Katılımcı dağılımı')
        table = document.add_table(rows=3, cols=2)
        table.rows[0].cells[0].text = 'Grup'
        table.rows[0].cells[1].text = 'Sayı'
        for cell in table.rows[0].cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True
        merged = table.cell(1, 0).merge(table.cell(2, 0))
        merged.text = 'Gezgin'
        table.cell(1, 1).text = '10'
        table.cell(2, 1).text = '20'
        table.columns[0].width = Cm(4)
        table.columns[1].width = Cm(8)

        png = base64.b64decode(
            'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9WlV4ssAAAAASUVORK5CYII='
        )
        document.add_picture(io.BytesIO(png), width=Cm(1))
        document.add_paragraph('Şekil 2: Örnek görsel')

        document.add_heading('Bulgular / Findings', level=1)
        document.add_paragraph('Bulgular bölümünün içeriğidir.')
        document.add_heading('Kaynakça', level=1)
        document.add_paragraph('Yılmaz, A. (2026). Örnek kaynak.')

        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def test_docx_import_extracts_structured_article_content(self):
        data, images = extract_form_data_from_docx(self.build_article_docx())

        self.assertEqual(data['cover']['tr_title'], 'Sürdürülebilir Turizmde Dijital Dönüşüm')
        self.assertEqual(data['cover']['en_title'], 'Digital Transformation in Sustainable Tourism')
        self.assertIn('Word içe aktarma', data['abstract']['tr_abs'])
        self.assertEqual(data['abstract']['tr_kw'], 'turizm; dijitalleşme')
        self.assertEqual(data['abstract']['en_kw'], 'tourism; digitalisation')
        self.assertEqual([section['name'] for section in data['sections']], [
            'Giriş / Introduction', 'Bulgular / Findings',
        ])
        self.assertIn('Araştırmanın başlangıç', data['sections'][0]['content'])

        table_item = next(item for item in data['figtables'] if item['type'] == 'table')
        self.assertEqual(table_item['number'], '1')
        self.assertEqual(table_item['tr_cap'], 'Katılımcı dağılımı')
        self.assertEqual(table_item['after_para'], 'Araştırmanın başlangıç paragrafıdır.')
        self.assertEqual(table_item['tbl_model']['rows'][1][0]['rowspan'], 2)
        self.assertTrue(table_item['tbl_model']['rows'][0][0]['bold'])

        figure_item = next(item for item in data['figtables'] if item['type'] == 'figure')
        self.assertEqual(figure_item['number'], '2')
        self.assertEqual(figure_item['tr_cap'], 'Örnek görsel')
        self.assertEqual(figure_item['import_image_index'], 0)
        self.assertEqual(len(images), 1)
        self.assertTrue(images[0]['mimetype'].startswith('image/'))
        self.assertEqual(data['references'], 'Yılmaz, A. (2026). Örnek kaynak.')
        self.assertEqual(data['extra'], {'ack': '', 'contrib': '', 'conflict': ''})
        self.assertEqual(data['stats']['tables'], 1)
        self.assertEqual(data['stats']['figures'], 1)

    def test_docx_import_preserves_optional_endmatter_only_when_present(self):
        document = Document()
        document.add_paragraph('English Manuscript', style='Title')
        document.add_paragraph('English Manuscript', style='Subtitle')
        document.add_heading('Abstract', level=1)
        document.add_paragraph('This abstract verifies optional end matter extraction.')
        document.add_paragraph('Keywords: tourism; review')
        document.add_heading('Conclusion', level=1)
        document.add_paragraph('Conclusion text.')
        document.add_heading('Author Contributions', level=1)
        document.add_paragraph('Both authors contributed to the study design and writing.')
        document.add_heading('Conflict of Interest', level=1)
        document.add_paragraph('The authors declare no conflict of interest.')
        document.add_heading('References', level=1)
        document.add_paragraph('Example, A. (2026). Reference.')

        output = io.BytesIO()
        document.save(output)
        data, _ = extract_form_data_from_docx(output.getvalue())

        self.assertEqual(
            data['extra']['contrib'],
            'Both authors contributed to the study design and writing.',
        )
        self.assertEqual(
            data['extra']['conflict'],
            'The authors declare no conflict of interest.',
        )
        self.assertEqual(data['references'], 'Example, A. (2026). Reference.')
        self.assertEqual([section['name'] for section in data['sections']], ['Conclusion'])

    def test_docx_import_builds_dynamic_sections_from_custom_headings(self):
        document = Document()
        document.add_paragraph('Özel Bölümlü Turizm Araştırması', style='Title')
        document.add_paragraph('Tourism Research with Custom Sections', style='Subtitle')
        document.add_heading('Özet', level=1)
        document.add_paragraph('Özet metni.')
        document.add_paragraph('Anahtar Kelimeler: turizm; deneyim')

        bold_heading = document.add_paragraph()
        bold_heading.add_run('1. Seyahat Motivasyonları').bold = True
        document.add_paragraph('Motivasyonlara ilişkin bölüm metni.')
        document.add_paragraph('2.1 Dijital Deneyim Tasarımı')
        document.add_paragraph('Alt bölümün içerik paragrafı.')
        document.add_heading('3. Özgün Öneriler', level=1)
        document.add_paragraph('Öneriler bölümünün içeriği.')
        document.add_heading('4. Kaynakça', level=1)
        document.add_paragraph('Kaya, B. (2026). Dinamik bölüm örneği.')

        output = io.BytesIO()
        document.save(output)
        data, _ = extract_form_data_from_docx(output.getvalue())

        self.assertEqual([section['name'] for section in data['sections']], [
            'Seyahat Motivasyonları',
            'Dijital Deneyim Tasarımı',
            'Özgün Öneriler',
        ])
        self.assertEqual([section['level'] for section in data['sections']], ['1', '2', '1'])
        self.assertIn('Motivasyonlara ilişkin', data['sections'][0]['content'])
        self.assertIn('Alt bölümün içerik', data['sections'][1]['content'])
        self.assertEqual(data['references'], 'Kaya, B. (2026). Dinamik bölüm örneği.')

    def test_docx_import_reads_outline_numbering_and_abstract_layouts(self):
        document = Document()
        document.add_paragraph('Hiyerarşik Bölümlü Araştırma', style='Title')
        document.add_paragraph('Research with Hierarchical Sections', style='Subtitle')

        abstract_table = document.add_table(rows=2, cols=1)
        abstract_table.cell(0, 0).text = 'Özet'
        abstract_table.cell(1, 0).text = 'Tablo düzenindeki Türkçe özet metni.'
        document.add_paragraph('Anahtar Kelimeler: turizm; hiyerarşi')
        document.add_paragraph('Abstract — English abstract on the label line.')
        document.add_paragraph('Keywords: tourism; hierarchy')

        main_style = document.styles.add_style('JGTTR Main Custom', WD_STYLE_TYPE.PARAGRAPH)
        main_style.font.bold = True
        main_ppr = main_style.element.get_or_add_pPr()
        main_outline = OxmlElement('w:outlineLvl')
        main_outline.set(qn('w:val'), '0')
        main_ppr.append(main_outline)

        sub_style = document.styles.add_style('JGTTR Sub Custom', WD_STYLE_TYPE.PARAGRAPH)
        sub_style.font.bold = True
        sub_ppr = sub_style.element.get_or_add_pPr()
        sub_outline = OxmlElement('w:outlineLvl')
        sub_outline.set(qn('w:val'), '1')
        sub_ppr.append(sub_outline)

        document.add_paragraph('Özel Ana Bölüm', style=main_style)
        document.add_paragraph('Ana bölümün metni.')
        document.add_paragraph('Özel Alt Bölüm', style=sub_style)
        document.add_paragraph('Alt bölümün metni.')

        numbered_heading = document.add_paragraph()
        numbered_heading.add_run('Otomatik Numaralı Alt-Alt Bölüm').bold = True
        p_pr = numbered_heading._p.get_or_add_pPr()
        num_pr = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl')
        ilvl.set(qn('w:val'), '2')
        num_id = OxmlElement('w:numId')
        num_id.set(qn('w:val'), '1')
        num_pr.append(ilvl)
        num_pr.append(num_id)
        p_pr.append(num_pr)
        document.add_paragraph('Alt-alt bölümün metni.')

        output = io.BytesIO()
        document.save(output)
        data, _ = extract_form_data_from_docx(output.getvalue())

        self.assertEqual(data['abstract']['tr_abs'], 'Tablo düzenindeki Türkçe özet metni.')
        self.assertEqual(data['abstract']['tr_kw'], 'turizm; hiyerarşi')
        self.assertEqual(data['abstract']['en_abs'], 'English abstract on the label line.')
        self.assertEqual(data['abstract']['en_kw'], 'tourism; hierarchy')
        self.assertEqual([section['name'] for section in data['sections']], [
            'Özel Ana Bölüm', 'Özel Alt Bölüm', 'Otomatik Numaralı Alt-Alt Bölüm',
        ])
        self.assertEqual([section['level'] for section in data['sections']], ['1', '2', '3'])
        self.assertEqual(data['figtables'], [])

    def test_large_word_table_is_marked_for_page_splitting(self):
        document = Document()
        document.add_paragraph('Büyük Tablolu Araştırma', style='Title')
        document.add_paragraph('Research with a Large Table', style='Subtitle')
        document.add_paragraph('Özet: Kısa Türkçe özet.')
        document.add_paragraph('Anahtar Kelimeler: tablo; turizm')
        document.add_heading('Bulgular', level=1)
        document.add_paragraph('Tablodan önceki paragraf.')
        table = document.add_table(rows=21, cols=3)
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                cell.text = f'R{row_index + 1} C{col_index + 1}'

        output = io.BytesIO()
        document.save(output)
        data, _ = extract_form_data_from_docx(output.getvalue())

        table_item = next(item for item in data['figtables'] if item['type'] == 'table')
        self.assertTrue(table_item['split_table'])

    def test_body_table_with_abstract_named_cell_is_not_flattened(self):
        document = Document()
        document.add_paragraph('Özet Hücreli Gövde Tablosu', style='Title')
        document.add_paragraph('Body Table with an Abstract-named Cell', style='Subtitle')
        document.add_paragraph('Özet: Bu çalışma normal gövde tablolarının korunmasını sınar.')
        document.add_paragraph('Anahtar Kelimeler: tablo; özet')
        document.add_heading('Bulgular', level=1)
        document.add_paragraph('Tablo 1: Değerlendirme sonuçları')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Kategori'
        table.cell(0, 1).text = 'Değer'
        table.cell(1, 0).text = 'Özet'
        table.cell(1, 1).text = 'Başarılı'

        output = io.BytesIO()
        document.save(output)
        data, _ = extract_form_data_from_docx(output.getvalue())

        self.assertEqual(data['stats']['tables'], 1)
        self.assertEqual(data['figtables'][0]['tbl_model']['rows'][1][0]['text'], 'Özet')
        self.assertEqual([section['name'] for section in data['sections']], ['Bulgular'])

    def test_abstracts_are_read_from_large_word_layout_tables(self):
        document = Document()
        document.add_paragraph('Yerleşim Tablolu Turizm Araştırması', style='Title')
        document.add_paragraph('Tourism Research with a Layout Table', style='Subtitle')

        layout = document.add_table(rows=4, cols=2)
        layout.cell(0, 0).text = 'TÜRKÇE ÖZET'
        layout.cell(0, 1).text = 'Bu çalışma, büyük bir Word yerleşim tablosunda saklanan Türkçe özetin eksiksiz aktarılmasını sınamaktadır.'
        layout.cell(1, 0).text = 'Anahtar Sözcükler'
        layout.cell(1, 1).text = 'turizm; yerleşim; aktarım'
        layout.cell(2, 0).text = 'ENGLISH ABSTRACT'
        layout.cell(2, 1).text = 'This study verifies extraction of an English abstract stored inside a large Word layout table.'
        layout.cell(3, 0).text = 'Index Terms'
        layout.cell(3, 1).text = 'tourism; layout; import'

        document.add_heading('Giriş', level=1)
        document.add_paragraph('Makale gövdesi.')
        output = io.BytesIO()
        document.save(output)

        data, _ = extract_form_data_from_docx(output.getvalue())

        self.assertIn('büyük bir Word yerleşim tablosunda', data['abstract']['tr_abs'])
        self.assertEqual(data['abstract']['tr_kw'], 'turizm; yerleşim; aktarım')
        self.assertIn('stored inside a large Word layout table', data['abstract']['en_abs'])
        self.assertEqual(data['abstract']['en_kw'], 'tourism; layout; import')
        self.assertEqual(data['figtables'], [])

    def test_article_type_before_bilingual_abstract_table_stays_front_matter(self):
        document = Document()
        document.add_paragraph('Araştırma Makalesi – Research Article', style='Heading 1')
        document.add_paragraph('Macera Turizmi Araştırması', style='Title')
        document.add_paragraph('Adventure Tourism Research', style='Subtitle')
        layout = document.add_table(rows=4, cols=2)
        layout.cell(0, 0).text = 'ÖZET'
        layout.cell(0, 1).text = 'Bu çalışma uzun Türkçe özetlerin güvenli aktarımını doğrulamaktadır.'
        layout.cell(1, 0).text = 'Anahtar Kelimeler'
        layout.cell(1, 1).text = 'turizm; tablo'
        layout.cell(2, 0).text = 'ABSTRACT'
        layout.cell(2, 1).text = 'This study verifies safe extraction of a long English abstract.'
        layout.cell(3, 0).text = 'Keywords'
        layout.cell(3, 1).text = 'tourism; table'
        document.add_heading('Giriş', level=1)
        document.add_paragraph('Giriş metni.')

        output = io.BytesIO()
        document.save(output)
        data, _ = extract_form_data_from_docx(output.getvalue())

        self.assertEqual(data['cover']['article_type'], 'Araştırma Makalesi – Research Article')
        self.assertIn('uzun Türkçe özetlerin', data['abstract']['tr_abs'])
        self.assertIn('long English abstract', data['abstract']['en_abs'])
        self.assertEqual(data['stats']['tables'], 0)
        self.assertEqual([section['name'] for section in data['sections']], ['Giriş'])

    def test_narrative_table_reference_does_not_replace_caption(self):
        document = Document()
        document.add_paragraph('Tablo Başlığı Denetimi', style='Title')
        document.add_paragraph('Table Caption Audit', style='Subtitle')
        document.add_paragraph('Özet: Bu çalışma tablo başlıklarını sınar.')
        document.add_paragraph('Anahtar Kelimeler: tablo; başlık')
        document.add_heading('Bulgular', level=1)
        document.add_paragraph('Tablo 1: Kısa ve doğru başlık')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'A'
        table.cell(0, 1).text = 'B'
        table.cell(1, 0).text = '1'
        table.cell(1, 1).text = '2'
        document.add_paragraph(
            'Tablo 1’de görüldüğü üzere bu cümle bir açıklama paragrafıdır ve başlık değildir.'
        )

        output = io.BytesIO()
        document.save(output)
        data, _ = extract_form_data_from_docx(output.getvalue())

        self.assertEqual(data['figtables'][0]['tr_cap'], 'Kısa ve doğru başlık')
        self.assertIn('bir açıklama paragrafıdır', data['sections'][0]['content'])

    def test_abstracts_use_style_and_keyword_fallback_when_labels_are_missing(self):
        document = Document()
        document.add_paragraph('Etiketsiz Özetli Araştırma', style='Title')
        document.add_paragraph('Research with Unlabelled Abstracts', style='Subtitle')

        abstract_style = document.styles.add_style('Türkçe Özet Metni', WD_STYLE_TYPE.PARAGRAPH)
        document.add_paragraph(
            'Bu çalışma Word belgesinde ayrı bir özet etiketi bulunmadığında paragraf stilinden Türkçe içeriği güvenilir biçimde belirlemeyi amaçlamaktadır.',
            style=abstract_style,
        )
        document.add_paragraph('Anahtar Kelimeler: turizm; word; özet')
        document.add_paragraph(
            'This study verifies the fallback that identifies an English abstract even when its separate heading is missing from the Word document.'
        )
        document.add_paragraph('Keywords: tourism; word; abstract')
        document.add_heading('Yöntem', level=1)
        document.add_paragraph('Yöntem bölümü.')
        output = io.BytesIO()
        document.save(output)

        data, _ = extract_form_data_from_docx(output.getvalue())

        self.assertIn('paragraf stilinden Türkçe içeriği', data['abstract']['tr_abs'])
        self.assertIn('separate heading is missing', data['abstract']['en_abs'])
        self.assertEqual(data['abstract']['tr_kw'], 'turizm; word; özet')
        self.assertEqual(data['abstract']['en_kw'], 'tourism; word; abstract')

    def test_structured_abstract_subheadings_are_not_body_sections(self):
        document = Document()
        document.add_paragraph('Yapılandırılmış Özet Araştırması', style='Title')
        document.add_paragraph('Structured Abstract Research', style='Subtitle')
        document.add_heading('Özet', level=1)
        purpose = document.add_paragraph()
        purpose.add_run('Amaç:').bold = True
        purpose.add_run(' Çalışmanın amacı özet ayrıştırmasını sınamaktır.')
        method = document.add_paragraph()
        method.add_run('Yöntem:').bold = True
        method.add_run(' Word belgesi programatik olarak incelenmiştir.')
        document.add_paragraph('Anahtar Kelimeler: özet; yöntem')
        document.add_heading('Giriş', level=1)
        document.add_paragraph('Giriş metni.')
        output = io.BytesIO()
        document.save(output)

        data, _ = extract_form_data_from_docx(output.getvalue())

        self.assertIn('Amaç:', data['abstract']['tr_abs'])
        self.assertIn('Yöntem:', data['abstract']['tr_abs'])
        self.assertEqual([section['name'] for section in data['sections']], ['Giriş'])

    def test_abstracts_inside_text_boxes_and_content_controls_are_recovered(self):
        document = Document()
        document.add_paragraph('Metin Kutulu Özet Araştırması', style='Title')
        document.add_paragraph('Research with Text-box Abstracts', style='Subtitle')

        def xml_paragraph(text):
            paragraph = OxmlElement('w:p')
            run = OxmlElement('w:r')
            node = OxmlElement('w:t')
            node.text = text
            run.append(node)
            paragraph.append(run)
            return paragraph

        outer = document.add_paragraph()
        text_box = OxmlElement('w:txbxContent')
        text_box.append(xml_paragraph('ÖZET: Bu özet Word metin kutusunun içinde saklanmaktadır.'))
        text_box.append(xml_paragraph('Anahtar Kelimeler: turizm; metin kutusu'))
        outer._p.append(text_box)

        content_control = OxmlElement('w:sdt')
        content_control.append(OxmlElement('w:sdtPr'))
        content = OxmlElement('w:sdtContent')
        content.append(xml_paragraph('ABSTRACT: This abstract is stored inside a Word content control.'))
        content.append(xml_paragraph('Keywords: tourism; content control'))
        content_control.append(content)
        document.element.body.insert(-1, content_control)

        document.add_heading('Giriş', level=1)
        document.add_paragraph('Gövde metni.')
        output = io.BytesIO()
        document.save(output)

        data, _ = extract_form_data_from_docx(output.getvalue())

        self.assertIn('Word metin kutusunun içinde', data['abstract']['tr_abs'])
        self.assertEqual(data['abstract']['tr_kw'], 'turizm; metin kutusu')
        self.assertIn('Word content control', data['abstract']['en_abs'])
        self.assertEqual(data['abstract']['en_kw'], 'tourism; content control')
