"""Editor feedback: short titles, automatic page ranges, bilingual panels and placements."""
import io
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_export import generate_docx_from_form
from formatter import extract_form_data_from_docx, generate_latex_from_form
from journal_templates import TEMPLATES, normalize_settings
from page_furniture import LAST_PAGE, article_values, citation_text
from test_formatter_templates import sample_article


PANEL = next(item['settings'] for item in TEMPLATES if item['id'] == 'bilingual_panel')


def docx(data, settings):
    return Document(io.BytesIO(generate_docx_from_form(data, {}, settings)))


def all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return '\n'.join(parts)


def footer_paragraphs(doc):
    return doc.sections[0].first_page_footer.paragraphs


class ShortTitleTests(unittest.TestCase):
    def test_short_title_feeds_running_headers(self):
        data = sample_article()
        data['cover']['short_title'] = 'Açık İş Akışları'
        self.assertEqual(article_values(data, normalize_settings(PANEL))['kisa_baslik'], 'Açık İş Akışları')
        tex = generate_latex_from_form(data, {}, PANEL)
        running = tex[tex.index('% Custom running headers and footers'):tex.index('% ── Caption format ──')]
        self.assertIn('Açık İş Akışları', running)
        self.assertNotIn(data['cover']['tr_title'], running)
        header = docx(data, PANEL).sections[1].header.tables[0].cell(0, 0).text
        self.assertEqual(header, 'Açık İş Akışları')

    def test_empty_short_title_keeps_previous_behaviour(self):
        data = sample_article()
        self.assertEqual(article_values(data, normalize_settings(PANEL))['kisa_baslik'], data['cover']['tr_title'])


class AutomaticPageRangeTests(unittest.TestCase):
    def test_empty_end_page_uses_last_page_in_latex(self):
        data = sample_article()
        data['cover']['end_page'] = ''
        tex = generate_latex_from_form(data, {}, {})
        self.assertIn(r'\newcommand{\JGTTRendpage}{\pageref*{LastPage}}', tex)
        self.assertIn(r'89--\pageref*{LastPage}', tex)
        self.assertNotIn('xxx', tex[tex.index(r'\newcommand{\JGTTRapacitation}'):][:400])
        panel = generate_latex_from_form({**data, 'cover': {**data['cover']}}, {}, {**PANEL, 'header_left': '{sayfa_araligi}'})
        self.assertIn(r'89–\pageref*{LastPage}', panel)

    def test_explicit_end_page_is_kept(self):
        tex = generate_latex_from_form(sample_article(), {}, {})
        self.assertIn(r'\newcommand{\JGTTRendpage}{104}', tex)
        self.assertNotIn(r'\pageref*{LastPage}', tex)

    def test_word_uses_pageref_to_document_end(self):
        data = sample_article()
        data['cover']['end_page'] = ''
        settings = {'citation_position': 'footer', 'header_mode': 'same', 'header_center': '{sayfa_araligi}'}
        raw = generate_docx_from_form(data, {}, settings)
        doc = Document(io.BytesIO(raw))
        self.assertIn('w:name="aiditor_last_page"', doc.element.xml)
        self.assertIn('PAGEREF aiditor_last_page', doc.sections[0].first_page_footer._element.xml)
        self.assertIn('PAGEREF aiditor_last_page', doc.sections[1].header._element.xml)
        self.assertIn('updateFields', doc.settings.element.xml)
        self.assertNotIn(LAST_PAGE, doc.element.xml)

    def test_word_without_auto_end_page_needs_no_field_update(self):
        doc = docx(sample_article(), {})
        self.assertNotIn('aiditor_last_page', doc.element.xml)
        self.assertNotIn('updateFields', doc.settings.element.xml)

    def test_short_author_does_not_repeat_year(self):
        data = sample_article()
        data['cover']['author_short'] = 'Örnek & Smith, 2026'
        self.assertTrue(citation_text(data, normalize_settings({})).startswith('Örnek & Smith (2026).'))
        tex = generate_latex_from_form(data, {}, {})
        self.assertIn(r'\newcommand{\JGTTRauthorshort}{Örnek \& Smith}', tex)


class BilingualPanelInfoTests(unittest.TestCase):
    def setUp(self):
        self.data = sample_article()
        self.data['cover']['article_type'] = 'Araştırma Makalesi -- Research Article'
        self.data['cover']['revisions'] = '01.05.2026\n15.06.2026'

    def test_latex_panels_split_type_and_list_all_dates(self):
        tex = generate_latex_from_form(self.data, {}, PANEL)
        self.assertIn(r'\newcommand{\JGTTRarticletypetr}{Araştırma Makalesi}', tex)
        self.assertIn(r'\newcommand{\JGTTRarticletypeen}{Research Article}', tex)
        layout = tex[tex.index('% Template: bilingual_panel;'):]
        self.assertIn(r'\textbf{\JGTTRarticletypetr}', layout)
        self.assertIn(r'\textbf{\JGTTRarticletypeen}', layout)
        self.assertNotIn(r'\textbf{\JGTTRarticletype}', layout)
        for text in ('Geliş: 14.03.2026', 'Revizyon 1: 01.05.2026', 'Revizyon 2: 15.06.2026',
                     'Kabul: 22.07.2026', 'Yayın: 08.09.2026', 'Revised 2: 15.06.2026', 'Published: 08.09.2026'):
            self.assertIn(text, layout)
        self.assertLess(layout.index('Geliş:'), layout.index('Revizyon 1:'))
        self.assertLess(layout.index('Revizyon 2:'), layout.index('Kabul:'))

    def test_word_panels_split_type_and_list_all_dates(self):
        doc = docx(self.data, PANEL)
        turkish, english = doc.tables[1].cell(0, 0).text, doc.tables[2].cell(0, 0).text
        self.assertIn('Araştırma Makalesi', turkish)
        self.assertNotIn('Research Article', turkish)
        self.assertIn('Research Article', english)
        self.assertNotIn('Araştırma', english)
        self.assertIn('Revizyon 1: 01.05.2026\nRevizyon 2: 15.06.2026\nKabul: 22.07.2026\nYayın: 08.09.2026', turkish)
        self.assertIn('Revised 1: 01.05.2026', english)

    def test_single_revision_has_no_number(self):
        self.data['cover']['revisions'] = '01.05.2026'
        self.assertIn('Revizyon: 01.05.2026', docx(self.data, PANEL).tables[1].cell(0, 0).text)

    def test_other_templates_show_optional_revisions(self):
        for template in ('classic', 'contemporary', 'scholarly'):
            settings = next(item['settings'] for item in TEMPLATES if item['id'] == template)
            self.assertIn('01.05.2026', generate_latex_from_form(self.data, {}, settings), template)
            self.assertIn('01.05.2026', all_text(docx(self.data, settings)) + '\n'.join(
                p.text for p in footer_paragraphs(docx(self.data, settings))), template)


class PlacementOptionTests(unittest.TestCase):
    def test_settings_validate_new_choices(self):
        self.assertEqual(normalize_settings({})['citation_position'], 'auto')
        self.assertEqual(normalize_settings({})['footer_text_align'], 'left')
        with self.assertRaises(ValueError):
            normalize_settings({'citation_position': 'header'})
        with self.assertRaises(ValueError):
            normalize_settings({'footer_text_align': 'justify'})

    def test_licence_text_can_move_right(self):
        settings = {**PANEL, 'footer_text': 'CC BY 4.0 lisanslıdır.', 'footer_text_align': 'right'}
        tex = generate_latex_from_form(sample_article(), {}, settings)
        self.assertIn(r'{\raggedleft CC BY 4.0 lisanslıdır.\par}', tex)
        notes = [p for p in footer_paragraphs(docx(sample_article(), settings)) if p.text == 'CC BY 4.0 lisanslıdır.']
        self.assertEqual(notes[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        for template in ('classic', 'scholarly'):
            template_settings = {**next(item['settings'] for item in TEMPLATES if item['id'] == template),
                                 'footer_text': 'Lisans notu', 'footer_text_align': 'center'}
            self.assertIn(r'{\centering Lisans notu\par}', generate_latex_from_form(sample_article(), {}, template_settings))
            notes = [p for p in footer_paragraphs(docx(sample_article(), template_settings)) if p.text == 'Lisans notu']
            self.assertEqual(notes[0].alignment, WD_ALIGN_PARAGRAPH.CENTER, template)

    def test_default_panel_has_no_citation(self):
        tex = generate_latex_from_form(sample_article(), {}, PANEL)
        self.assertNotIn('Önerilen Atıf', tex[tex.index('% Template: bilingual_panel;'):])
        self.assertNotIn('Önerilen Atıf', all_text(docx(sample_article(), PANEL)))

    def test_panel_citation_between_authors_and_abstracts(self):
        settings = {**PANEL, 'citation_position': 'below_authors'}
        layout = generate_latex_from_form(sample_article(), {}, settings)
        layout = layout[layout.index('% Template: bilingual_panel;'):]
        self.assertLess(layout.index(r'\JGTTRaffiliations'), layout.index('Önerilen Atıf'))
        self.assertLess(layout.index('Önerilen Atıf'), layout.index(r'\colorbox{JGTTRgray!35}'))
        doc = docx(sample_article(), settings)
        paragraphs = [p.text for p in doc.paragraphs]
        citation = next(i for i, text in enumerate(paragraphs) if text.startswith('Önerilen Atıf'))
        self.assertIn('Örnek & Smith (2026).', paragraphs[citation])
        self.assertTrue(any('Örnek Üniversitesi' in text for text in paragraphs[:citation]))

    def test_citation_footer_and_none_for_standard_templates(self):
        classic = next(item['settings'] for item in TEMPLATES if item['id'] == 'classic')
        footer = generate_latex_from_form(sample_article(), {}, {**classic, 'citation_position': 'footer', 'footer_layout': 'minimal'})
        self.assertIn('Önerilen Atıf / Suggested Citation', footer)
        hidden = generate_latex_from_form(sample_article(), {}, {**classic, 'citation_position': 'none'})
        self.assertNotIn('Önerilen Atıf', hidden)
        doc = docx(sample_article(), {**classic, 'citation_position': 'none'})
        self.assertFalse(any('Önerilen Atıf' in p.text for p in footer_paragraphs(doc)))
        for template in ('classic', 'contemporary', 'centered', 'minimal'):
            settings = {**next(item['settings'] for item in TEMPLATES if item['id'] == template), 'citation_position': 'below_authors'}
            tex = generate_latex_from_form(sample_article(), {}, settings)
            cover = tex[tex.index(r'\newcommand{\JGTTRfirstpage}'):]
            self.assertLess(cover.index('#1'), cover.index('Önerilen Atıf'), template)
            self.assertNotIn('Önerilen Atıf', tex[tex.index(r'\newcommand{\JGTTRfirstpagefooter}'):tex.index(r'\newcommand{\JGTTRfirstpage}')])
        doc = docx(sample_article(), {**classic, 'citation_position': 'below_authors'})
        self.assertTrue(any(p.text.startswith('Önerilen Atıf') for p in doc.paragraphs))


class ImportedTitleLanguageTests(unittest.TestCase):
    def build(self, *titles, title_style_index=0):
        document = Document()
        for index, title in enumerate(titles):
            document.add_paragraph(title, style='Title' if index == title_style_index else None)
        document.add_paragraph('Öz')
        document.add_paragraph('Bu çalışma örnek bir özet içermektedir ve test amacıyla yazılmıştır.')
        document.add_heading('Giriş', level=1)
        document.add_paragraph('Giriş metni.')
        stream = io.BytesIO(); document.save(stream)
        return extract_form_data_from_docx(stream.getvalue())[0]['cover']

    def test_english_first_title_goes_to_english_field(self):
        cover = self.build('Digital Promotion and Visitor Experience in Rural Tourism',
                           'Kırsal Turizmde Dijital Tanıtım ve Ziyaretçi Deneyimi')
        self.assertEqual(cover['tr_title'], 'Kırsal Turizmde Dijital Tanıtım ve Ziyaretçi Deneyimi')
        self.assertEqual(cover['en_title'], 'Digital Promotion and Visitor Experience in Rural Tourism')

    def test_turkish_first_order_is_unchanged(self):
        cover = self.build('Kırsal Turizmde Dijital Tanıtım ve Ziyaretçi Deneyimi',
                           'Digital Promotion and Visitor Experience in Rural Tourism')
        self.assertEqual(cover['tr_title'], 'Kırsal Turizmde Dijital Tanıtım ve Ziyaretçi Deneyimi')

    def test_only_english_title_is_not_put_in_turkish_field(self):
        cover = self.build('Digital Promotion and Visitor Experience in Rural Tourism')
        self.assertEqual(cover['tr_title'], '')
        self.assertEqual(cover['en_title'], 'Digital Promotion and Visitor Experience in Rural Tourism')


if __name__ == '__main__':
    unittest.main()
