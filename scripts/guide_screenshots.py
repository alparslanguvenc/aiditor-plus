"""Regenerate the user-guide screenshots in docs/images/kilavuz with disposable storage.

Uses fictional sample data only (journal, authors, logo, article). Requires Playwright Chromium:
    python -m playwright install chromium
    python scripts/guide_screenshots.py
"""
import argparse
from pathlib import Path
import sys
import tempfile

from docx import Document
from PIL import Image, ImageChops, ImageDraw, ImageFont
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / 'tests')]
from browser_workflow import Server  # noqa: E402

HIDE = 'html{scroll-behavior:auto!important} #app-toast{display:none!important}'


def font(size):
    for name in ['/System/Library/Fonts/Supplemental/Georgia Bold.ttf', 'DejaVuSans-Bold.ttf']:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            pass
    return ImageFont.load_default()


def sample_files(work):
    logo = Image.new('RGB', (360, 360), 'white'); draw = ImageDraw.Draw(logo)
    draw.ellipse((20, 20, 340, 340), outline='#1f5f66', width=14)
    draw.text((180, 182), 'ÖSA', anchor='mm', fill='#1f5f66', font=font(110))
    logo.save(work / 'ornek_logo.png')
    chart = Image.new('RGB', (1200, 700), 'white'); draw = ImageDraw.Draw(chart)
    draw.line((120, 600, 1120, 600), fill='#333', width=3); draw.line((120, 80, 120, 600), fill='#333', width=3)
    for i, (label, value) in enumerate([('Sosyal medya', 18), ('Yorumlar', 11), ('Rezervasyon', 7)]):
        x = 220 + i * 300
        draw.rectangle((x, 600 - value * 26, x + 160, 600), fill=['#1f5f66', '#5b9aa0', '#a9cfd2'][i])
        draw.text((x + 80, 635), label, anchor='mm', fill='#222', font=font(34))
        draw.text((x + 80, 572 - value * 26), str(value), anchor='mm', fill='#222', font=font(34))
    chart.save(work / 'ornek_sekil.png')

    doc = Document()
    doc.add_paragraph('Kırsal Turizmde Dijital Tanıtım ve Ziyaretçi Deneyimi', style='Title')
    doc.add_paragraph('Digital Promotion and Visitor Experience in Rural Tourism')
    doc.add_paragraph('Öz').runs[0].bold = True
    doc.add_paragraph('Bu çalışma, kırsal turizm işletmelerinin dijital tanıtım araçlarını nasıl kullandığını ve bu araçların ziyaretçi deneyimine etkisini incelemektedir. Örnek bir bölgede faaliyet gösteren işletmelerle yarı yapılandırılmış görüşmeler yapılmış ve veriler tematik analizle değerlendirilmiştir. Bulgular, sosyal medya içeriklerinin beklenti oluşturmada belirleyici olduğunu, ancak tutarlı bir dijital stratejinin çoğu işletmede bulunmadığını göstermektedir.')
    doc.add_paragraph('Anahtar Kelimeler: kırsal turizm; dijital pazarlama; ziyaretçi deneyimi')
    doc.add_paragraph('Abstract').runs[0].bold = True
    doc.add_paragraph('This study examines how rural tourism businesses use digital promotion tools and how these tools shape visitor experience. Semi-structured interviews were conducted with businesses operating in a sample region and the data were evaluated through thematic analysis. The findings show that social media content is decisive in forming expectations, while most businesses lack a consistent digital strategy.')
    doc.add_paragraph('Keywords: rural tourism; digital marketing; visitor experience')
    doc.add_heading('Giriş', level=1)
    doc.add_paragraph('Kırsal turizm, yerel kalkınmanın önemli araçlarından biri olarak görülmektedir (Yılmaz, 2020). Son yıllarda dijital kanallar, küçük ölçekli işletmelerin ziyaretçilere ulaşmasında belirleyici hâle gelmiştir.')
    doc.add_paragraph('Kaya ve Demir (2021), çevrim içi görünürlüğün ziyaretçi beklentisini doğrudan etkilediğini belirtmektedir. Bu çalışma söz konusu ilişkiyi işletme bakış açısından ele almaktadır.')
    doc.add_heading('Yöntem', level=1)
    doc.add_paragraph('Araştırmada nitel yöntem benimsenmiş, 12 işletme temsilcisiyle görüşme yapılmıştır. Görüşmeler ortalama 45 dakika sürmüştür.')
    doc.add_heading('Bulgular', level=1)
    doc.add_paragraph('Görüşmelerden elde edilen temalar Tablo 1’de özetlenmektedir.')
    table = doc.add_table(rows=4, cols=3); table.style = 'Table Grid'
    rows = [('Tema', 'Kod sayısı', 'Örnek ifade'), ('Sosyal medya', 18, 'Fotoğraflar misafirleri getiriyor'),
            ('Çevrim içi yorumlar', 11, 'Yorumları her gün okuyoruz'), ('Rezervasyon', 7, 'Telefon hâlâ daha yaygın')]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)
    doc.add_paragraph('Smith et al. (2022) tarafından önerilen çerçeve bu bulgularla örtüşmektedir.')
    doc.add_heading('Sonuç', level=1)
    doc.add_paragraph('Kırsal turizm işletmeleri için sade ama tutarlı bir dijital tanıtım planı, ziyaretçi deneyimini güçlendirmektedir.')
    doc.add_heading('Kaynakça', level=1)
    for ref in ['Kaya, B., & Demir, C. (2021). Çevrim içi görünürlük ve ziyaretçi beklentisi. Örnek Turizm Araştırmaları, 5(1), 20–34. https://doi.org/10.1234/example.2021.5',
                'Smith, J., Brown, L., & Green, P. (2022). Digital tools for small tourism firms. Example Journal of Tourism, 12(3), 101–118. https://doi.org/10.1234/example.2022.12',
                'Yılmaz, A. (2020). Kırsal kalkınma ve turizm. Örnek Yayınevi.']:
        doc.add_paragraph(ref)
    doc.save(work / 'ornek_makale.docx')


def wait_ready(page):
    page.wait_for_function('window.journalWorkspace?.ready && window.articleLibrary?.ready')


def wait_saved(page):
    page.wait_for_function('window.articleLibrary?.currentId && !articleLibrary.dirty && document.getElementById("autosave-status").dataset.state === "saved"', timeout=20000)


def wait_preset(page):
    page.wait_for_function('window.journalWorkspace?.ready && !journalWorkspace.dirty && document.getElementById("journal-save-status").dataset.state === "saved"', timeout=20000)


def capture(url, work, out):
    def element(locator, name):
        locator.scroll_into_view_if_needed(); locator.page.wait_for_timeout(250)
        locator.screenshot(path=str(out / name))

    def region(page, name, selectors, pad=12):
        boxes = [page.locator(s).first.bounding_box() for s in selectors]
        sx, sy = page.evaluate('[scrollX, scrollY]')
        x0 = max(0, min(b['x'] for b in boxes) + sx - pad); y0 = max(0, min(b['y'] for b in boxes) + sy - pad)
        x1 = max(b['x'] + b['width'] for b in boxes) + sx + pad; y1 = max(b['y'] + b['height'] for b in boxes) + sy + pad
        page.screenshot(path=str(out / name), full_page=True, clip=dict(x=x0, y=y0, width=x1 - x0, height=y1 - y0))

    with sync_playwright() as p:
        browser = p.chromium.launch(args=['--lang=tr-TR'])
        page = browser.new_page(viewport={'width': 1440, 'height': 900}, device_scale_factor=2, locale='tr-TR')
        errors = []; page.on('pageerror', lambda e: errors.append(str(e))); page.on('dialog', lambda d: d.accept())
        page.goto(url, wait_until='networkidle'); page.add_style_tag(content=HIDE)
        page.screenshot(path=str(out / '01-giris-ekrani.png'))
        page.locator('#register-tab').click()
        page.locator('#auth-display-name').fill('Örnek Sosyal Araştırmalar Dergisi')
        page.locator('#auth-username').fill('ornek_dergi')
        page.locator('#auth-password').fill('ornek-parola-2026')
        element(page.locator('.auth-card'), '02-hesap-olustur.png')
        page.locator('#auth-submit').click(); wait_ready(page); page.wait_for_timeout(600)
        page.add_style_tag(content=HIDE)

        # Journal design and settings
        page.locator('#journal-tab').click()
        page.locator('[data-template-id="contemporary"]').click(); wait_preset(page)
        page.locator('#js-name-en').fill('Journal of Sample Social Research')
        page.locator('#js-issn-online').fill('1234-5678')
        page.locator('#js-url').fill('https://example.org/osad')
        page.locator('#logo-inp').set_input_files(str(work / 'ornek_logo.png'))
        page.locator('#js-footer-text').fill('Bu makale CC BY 4.0 lisansı ile yayımlanmaktadır.')
        wait_preset(page); page.wait_for_timeout(500)
        page.evaluate('window.scrollTo({top:0,behavior:"instant"})'); page.wait_for_timeout(400)
        page.screenshot(path=str(out / '03-calisma-alani.png'))
        region(page, '04-sablon-galerisi.png', ['.setup-heading', '#template-gallery', '.preview-note'])
        sections = page.locator('#journal-form .settings-section')
        region(page, '05-dergi-kimligi.png', ['#journal-form .settings-section >> nth=0', '.setup-aside .preview-card'], pad=4)
        element(sections.nth(1), '06-sayfa-tasarimi.png')
        element(sections.nth(2), '07-dipnot-lisans.png')
        page.locator('#js-header_mode').select_option('odd_even')
        page.locator('#js-header_rule').select_option('line')
        page.locator('#js-footer_mode').select_option('same')
        page.locator('#js-footer_center').fill('Sayfa {sayfa}')
        wait_preset(page); page.wait_for_timeout(400)
        element(page.locator('.running-settings'), '08-ust-alt-bilgiler.png')
        region(page, '09-atif-ve-kayit.png', ['#journal-form .settings-section >> nth=4', '.journal-savebar'], pad=2)
        element(page.locator('.setup-aside > div').nth(1), '10-ayar-yedekleme.png')

        # Article: Word import, metadata, authors, figures and outputs
        page.locator('#articles-tab').click(); page.wait_for_timeout(300)
        page.locator('#new-article').click(); page.wait_for_timeout(300)
        page.locator('#docx-file-input').set_input_files(str(work / 'ornek_makale.docx'))
        page.wait_for_function('!docxImporting && collectSections().length > 2', timeout=20000)
        page.wait_for_timeout(500); wait_saved(page)
        cover = {'#c-year': '2026', '#c-volume': '3', '#c-issue': '2', '#c-start': '45',
                 '#c-doi': '10.1234/osad.2026.012', '#c-article-type': 'Araştırma Makalesi / Research Article',
                 '#c-received': '12.03.2026', '#c-accepted': '04.07.2026', '#c-published': '30.09.2026',
                 '#c-author-short': 'Örnek & Smith', '#c-editor': 'Dr. Deniz Editör',
                 '#c-short-title': 'Kırsal Turizmde Dijital Tanıtım',
                 '#c-ethics': 'Bu çalışma için Örnek Üniversitesi Etik Kurulundan izin alınmıştır (Karar no: 2026/04).'}
        for selector, value in cover.items():
            page.locator(selector).fill(value)
        page.locator('#c-end').fill('')
        page.locator('#add-revision').click(); page.keyboard.type('20.05.2026')
        authors = [dict(title='Doç. Dr.', name='Deniz Örnek', affiliation='Örnek Üniversitesi, Turizm Fakültesi, Ankara',
                        orcid='0000-0002-1825-0097', email='deniz@example.org', corresponding=True),
                   dict(title='Dr.', name='Alex Smith', affiliation='Example University, School of Tourism, London',
                        orcid='', email='alex@example.org', corresponding=False)]
        while page.locator('#authors-list .author-row').count() < len(authors):
            page.evaluate('addAuthor()')
        for i, author in enumerate(authors):
            row = page.locator('#authors-list .author-row').nth(i)
            for field in ['title', 'name', 'affiliation', 'orcid', 'email']:
                row.locator(f'[data-f="{field}"]').fill(author[field])
            if author['corresponding']:
                row.locator('[data-f="corresponding"]').check()
        figure = page.evaluate('String(addFT({type:"figure"}))')
        block = page.locator(f'#ft-{figure}')
        block.locator('[data-f="tr_cap"]').fill('Temalara göre kod sayıları')
        block.locator('[data-f="en_cap"]').fill('Code counts by theme')
        page.locator(f'#finp-{figure}').set_input_files(str(work / 'ornek_sekil.png'))
        table = page.locator('#ft-list .ft-block').first
        table.locator('[data-f="tr_cap"]').fill('Görüşmelerden elde edilen temalar')
        table.locator('[data-f="en_cap"]').fill('Themes derived from the interviews')
        page.wait_for_timeout(500); wait_saved(page)
        region(page, '11-makale-arsivi.png', ['.article-library', '.docx-import-card'], pad=4)
        for card, name in [('b-cover', '12-kapak-bilgileri'), ('b-authors', '13-yazarlar'), ('b-abstract', '14-ozet'),
                           ('b-sections', '15-bolumler'), ('b-ft', '16-sekil-tablolar'), ('b-extra', '17-ek-bolumler'),
                           ('b-refs', '18-kaynakca')]:
            element(page.locator(f'.card:has(#{card})'), name + '.png')
        page.locator('#btn-docx').click()
        expect(page.locator('#result-title')).to_have_text('Tamamlandı — Word belgesi hazır', timeout=60000)
        page.wait_for_timeout(400)
        region(page, '19-word-ciktisi.png', ['#btn-docx', '#result-panel'], pad=8)
        page.locator('#btn-gen').click()
        expect(page.locator('#latex-next-steps')).to_be_visible(timeout=60000)
        page.wait_for_timeout(400)
        element(page.locator('#result-panel'), '20-latex-ciktisi.png')

        browser.close()
        if errors:
            raise RuntimeError(errors)


def optimise(out, max_width):
    for path in sorted(out.glob('*.png')):
        image = Image.open(path).convert('RGB')
        # Remove the empty band below short pages.
        background = Image.new('RGB', image.size, image.getpixel((image.width - 1, image.height - 1)))
        box = ImageChops.difference(image, background).getbbox()
        if box and image.height - box[3] > 40:
            image = image.crop((0, 0, image.width, box[3] + 40))
        if image.width > max_width:
            image = image.resize((max_width, round(image.height * max_width / image.width)), Image.LANCZOS)
        image.quantize(colors=256, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE).save(path, optimize=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, default=ROOT / 'docs' / 'images' / 'kilavuz')
    parser.add_argument('--max-width', type=int, default=1600)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='aiditor-guide-') as data:
        work = Path(data)
        sample_files(work)
        server = Server(data)
        try:
            capture(server.start(), work, args.output)
        finally:
            server.stop()
    optimise(args.output, args.max_width)
    print('Kılavuz görselleri:', args.output)


if __name__ == '__main__':
    main()
